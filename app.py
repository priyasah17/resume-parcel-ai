
import os
import json

import streamlit as st
import pandas as pd
from dotenv import load_dotenv
from groq import Groq
from pydantic import BaseModel, Field
from pypdf import PdfReader
from docx import Document

# ---------------- PAGE CONFIG ----------------

st.set_page_config(
    page_title="Resume Parcel AI",
    page_icon="📄",
    layout="wide",
)

load_dotenv()
API_KEY = os.getenv("GROQ_API_KEY")
MODEL = "openai/gpt-oss-120b"

# ---------------- DATA SCHEMAS ----------------

class JobD(BaseModel):
    role: str
    required_skills: list[str] = Field(default_factory=list)
    preferred_skills: list[str] = Field(default_factory=list)
    minimum_experience: float | None = None
    education_requirements: list[str] = Field(default_factory=list)
    responsibilities: list[str] = Field(default_factory=list)


class Experience(BaseModel):
    company: str | None = None
    role: str | None = None
    duration: str | None = None
    description: str | None = None
    skills_used: list[str] = Field(default_factory=list)


class Resume(BaseModel):
    name: str | None = None
    email: str | None = None
    phone: str | None = None
    total_experience_years: float | None = None
    skills: list[str] = Field(default_factory=list)
    experiences: list[Experience] = Field(default_factory=list)
    education: list[str] = Field(default_factory=list)
    projects: list[str] = Field(default_factory=list)
    certifications: list[str] = Field(default_factory=list)


class MatchResult(BaseModel):
    score: float
    details: dict


# ---------------- AI FUNCTIONS ----------------

def get_client():
    if not API_KEY:
        raise ValueError(
            "GROQ_API_KEY is missing. Check your .env file."
        )
    return Groq(api_key=API_KEY)


def call_json(prompt, system_prompt="You are a helpful AI assistant."):
    client = get_client()

    response = client.chat.completions.create(
        model=MODEL,
        messages=[
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": prompt},
        ],
        response_format={"type": "json_object"},
        temperature=0,
    )

    content = response.choices[0].message.content
    if not content:
        raise ValueError("The AI returned an empty response.")

    return json.loads(content)


def parse_job_description(description):
    schema = JobD.model_json_schema()

    prompt = f"""
Extract job information from the following job description.

Return actual extracted values matching this JSON schema:
{json.dumps(schema)}

Do not return the schema itself.
Do not invent information.
If experience is not specified, use null.
If a list has no information, use an empty list.

JOB DESCRIPTION:
{description}
"""

    data = call_json(prompt)
    return JobD(**data)


def read_resume(uploaded_file):
    filename = uploaded_file.name.lower()

    if filename.endswith(".pdf"):
        reader = PdfReader(uploaded_file)
        pages = [
            page.extract_text() or ""
            for page in reader.pages
        ]
        return "\n".join(pages)

    if filename.endswith(".docx"):
        document = Document(uploaded_file)
        parts = [
            p.text for p in document.paragraphs
            if p.text.strip()
        ]

        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        parts.append(cell.text)

        return "\n".join(parts)

    raise ValueError("Only PDF and DOCX files are supported.")


def parse_resume(resume_text):
    schema = Resume.model_json_schema()

    prompt = f"""
Extract candidate information from this resume.

Return JSON matching this schema:
{json.dumps(schema)}

Rules:
- Do not invent information.
- Use null for unavailable single values.
- Use empty lists for unavailable list values.
- Include internships in experiences.
- Extract skills from the whole resume, including projects.
- Preserve the candidate's actual details.

RESUME:
{resume_text[:25000]}
"""

    data = call_json(prompt, "You are an expert resume parser.")
    return Resume(**data)


def score_resume(job, resume):
    schema = MatchResult.model_json_schema()

    prompt = f"""
Compare this candidate with the job requirements.

JOB:
{job.model_dump_json(indent=2)}

RESUME:
{resume.model_dump_json(indent=2)}

Return JSON matching this schema:
{json.dumps(schema)}

Requirements:
- score must be a numeric match percentage from 0 to 100.
- details must be an object containing:
  candidate_name, matching_skills, missing_important_skills,
  experience_requirement_met, overall_match_percentage,
  final_verdict.
- Use only evidence from the supplied job description and resume.
- Do not invent candidate qualifications.
- If the job has no experience requirement, use null for
  experience_requirement_met.
- Assess skills and qualifications fairly.
"""

    data = call_json(prompt, "You are a careful technical recruiter.")
    result = MatchResult(**data)

    if not 0 <= result.score <= 100:
        raise ValueError("AI returned a score outside 0–100.")

    return result


# ---------------- STREAMLIT INTERFACE ----------------

st.title("📄 Resume Parcel AI")
st.subheader("AI-Powered Resume Screening & Job Matching")

st.write(
    "Upload resumes, enter a job description, and compare "
    "candidate skills using AI."
)

st.info(
    "AI-generated scores are estimates to support human review, "
    "not automatic hiring decisions."
)

job_description = st.text_area(
    "📋 Job Description",
    height=260,
    placeholder="Paste the complete technical job description here...",
)

uploaded_resumes = st.file_uploader(
    "📂 Upload candidate resumes",
    type=["pdf", "docx"],
    accept_multiple_files=True,
    help="Upload PDF or DOCX resumes. Avoid unnecessary personal data.",
)

if uploaded_resumes:
    st.write(f"**Files selected:** {len(uploaded_resumes)}")
    for file in uploaded_resumes:
        st.write(f"📄 {file.name}")

if st.button("🚀 Analyze Resumes", type="primary", use_container_width=True):
    if not job_description.strip():
        st.error("Please enter a job description.")
    elif not uploaded_resumes:
        st.error("Please upload at least one resume.")
    elif not API_KEY:
        st.error("GROQ_API_KEY is missing. Check your .env file.")
    else:
        results = []

        try:
            with st.spinner("Analyzing the job description..."):
                job = parse_job_description(job_description)

            st.success(f"Job description analyzed: {job.role}")

            progress = st.progress(0)
            status = st.empty()

            for index, file in enumerate(uploaded_resumes):
                status.write(
                    f"Processing {index + 1}/{len(uploaded_resumes)}: "
                    f"{file.name}"
                )

                try:
                    file.seek(0)
                    resume_text = read_resume(file)

                    if not resume_text.strip():
                        raise ValueError(
                            "No readable text found. The PDF may be scanned."
                        )

                    parsed_resume = parse_resume(resume_text)
                    result = score_resume(job, parsed_resume)

                    results.append({
                        "Candidate": (
                            parsed_resume.name or file.name
                        ),
                        "Resume file": file.name,
                        "Match score (%)": round(result.score, 1),
                        "Resume": parsed_resume,
                        "Details": result.details,
                    })

                except Exception as exc:
                    st.warning(f"Could not process {file.name}: {exc}")

                progress.progress((index + 1) / len(uploaded_resumes))

            status.empty()
            progress.empty()

            if results:
                results.sort(
                    key=lambda item: item["Match score (%)"],
                    reverse=True,
                )

                st.session_state["resume_results"] = results
                st.session_state["job_data"] = job.model_dump()
            else:
                st.error(
                    "No resumes were analyzed successfully. "
                    "Check the files and API configuration."
                )

        except Exception as exc:
            st.error(f"Analysis failed: {exc}")


# ---------------- DISPLAY RESULTS ----------------

if "resume_results" in st.session_state:
    results = st.session_state["resume_results"]

    st.divider()
    st.header("📊 Candidate Rankings")

    col1, col2, col3 = st.columns(3)

    col1.metric("Resumes analyzed", len(results))
    col2.metric("Highest match", f"{results[0]['Match score (%)']}%")
    col3.metric(
        "Average match",
        f"{sum(r['Match score (%)'] for r in results) / len(results):.1f}%",
    )

    table_data = [
        {
            "Rank": i + 1,
            "Candidate": r["Candidate"],
            "Resume file": r["Resume file"],
            "Match score (%)": r["Match score (%)"],
        }
        for i, r in enumerate(results)
    ]

    st.dataframe(
        pd.DataFrame(table_data),
        use_container_width=True,
        hide_index=True,
    )

    csv_data = pd.DataFrame(table_data).to_csv(index=False).encode("utf-8")
    st.download_button(
        "⬇️ Download rankings as CSV",
        data=csv_data,
        file_name="resume_rankings.csv",
        mime="text/csv",
    )

    st.subheader("🔎 Candidate Details")

    for index, item in enumerate(results):
        details = item["Details"]
        resume = item["Resume"]

        with st.expander(
            f"{index + 1}. {item['Candidate']} — "
            f"{item['Match score (%)']}%"
        ):
            st.write(f"**Resume file:** {item['Resume file']}")
            st.write(f"**Email:** {resume.email or 'Not provided'}")
            st.write(f"**Experience:** {resume.total_experience_years if resume.total_experience_years is not None else 'Not specified'} years")

            matching = details.get(
                "matching_skills",
                details.get("matching_skills_found", []),
            )
            missing = details.get(
                "missing_important_skills",
                details.get("missing_skills", []),
            )
            verdict = details.get(
                "final_verdict",
                details.get("verdict", "No verdict returned."),
            )

            left, right = st.columns(2)

            with left:
                st.markdown("**✅ Matching skills**")
                if matching:
                    for skill in matching:
                        st.write(f"• {skill}")
                else:
                    st.write("No matching skills returned.")

            with right:
                st.markdown("**📌 Skills to develop / not found**")
                if missing:
                    for skill in missing:
                        st.write(f"• {skill}")
                else:
                    st.write("No missing skills returned.")

            st.markdown("**📝 AI assessment**")
            st.write(verdict)

            with st.expander("View extracted resume information"):
                st.json(resume.model_dump())

            with st.expander("View complete AI scoring details"):
                st.json(details)
